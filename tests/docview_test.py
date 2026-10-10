"""Split view (server/docview.py): Word documents as escaped HTML, workbooks and CSVs as grids with formulas and
merges, PDFs passed through, and the endpoint's path checks."""
import shutil
import unittest
import uuid
from unittest import mock

from smoke_test import client, shellb  # noqa: F401

import docview
from tools import WORKSPACES


class DocView(unittest.TestCase):
    def setUp(self):
        self.conv = f"conv-{uuid.uuid4().hex[:12]}"
        self.dir = WORKSPACES / self.conv
        self.dir.mkdir(parents=True)

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)
        docview._cache.clear()

    def get(self, path):
        return client.get(f"/api/docview/{self.conv}/{path}")

    def test_docx_is_escaped_html(self):
        import docx
        d = docx.Document()
        d.add_heading("Report", 0)
        p = d.add_paragraph("Growth ")
        p.add_run("12%").bold = True
        p.add_run(' <img src=x onerror="alert(1)">')
        d.add_paragraph("one", style="List Bullet")
        d.add_paragraph("two", style="List Bullet")
        t = d.add_table(rows=1, cols=3)
        t.cell(0, 0).merge(t.cell(0, 1)).text = "wide"
        d.save(str(self.dir / "r.docx"))
        r = self.get("r.docx")
        self.assertEqual(r.status_code, 200)
        h = r.json()["html"]
        self.assertEqual(r.json()["kind"], "doc")
        self.assertIn("<h1>Report</h1>", h)
        self.assertIn("<strong>12%</strong>", h)
        self.assertIn("&lt;img src=x onerror=&quot;alert(1)&quot;&gt;", h)
        self.assertNotIn("<img", h)
        self.assertIn("<ul><li>one</li><li>two</li></ul>", h)
        self.assertIn("colspan=2", h)

    def test_xlsx_grid(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(["Region", "Q1", "Share"])
        ws["A1"].font = openpyxl.styles.Font(bold=True)
        ws.append(["North", 10.5, 0.25])
        ws["C2"].number_format = "0.0%"
        ws["B3"] = "=B2*2"
        ws.merge_cells("A5:C5")
        wb.create_sheet("Hidden").sheet_state = "hidden"
        wb.save(str(self.dir / "s.xlsx"))
        # openpyxl saves no formula results; without LibreOffice the grid says so instead of failing
        with mock.patch("doc_tools.soffice_convert", side_effect=RuntimeError("no LibreOffice")):
            r = self.get("s.xlsx").json()
        self.assertEqual(r["kind"], "sheet")
        s = r["sheets"][0]
        self.assertEqual(s["rows"][1], ["North", "10.5", "25.0%"])
        self.assertEqual(s["formulas"], {"2,1": "=B2*2"})
        self.assertEqual(s["bold"], [[0, 0]])
        self.assertEqual(s["merges"], [[4, 0, 4, 2]])
        self.assertTrue(r["sheets"][1]["hidden"])
        self.assertIn("note", r)

    def test_csv_sniffs_delimiter(self):
        (self.dir / "d.csv").write_text("a;b\n1;2\n3\n")
        s = self.get("d.csv").json()["sheets"][0]
        self.assertEqual(s["rows"], [["a", "b"], ["1", "2"], ["3", ""]])

    def test_pdf_and_text(self):
        (self.dir / "x.pdf").write_bytes(b"%PDF-1.4\n")
        self.assertEqual(self.get("x.pdf").json(), {**self.get("x.pdf").json(), "kind": "pdf", "url": f"/api/workspace/{self.conv}/x.pdf"})
        (self.dir / "n.md").write_text("# Hi")
        self.assertEqual(self.get("n.md").json()["kind"], "markdown")

    def test_refuses_other_files_and_paths(self):
        (self.dir / "a.bin").write_bytes(b"\0")
        self.assertEqual(self.get("a.bin").status_code, 415)
        self.assertEqual(self.get("missing.docx").status_code, 404)
        self.assertIn(self.get("..%2F..%2Fshellb.db").status_code, (400, 404))
        self.assertEqual(client.get(f"/api/docview/{self.conv}/x.pdf?pdf=1").status_code, 404)


if __name__ == "__main__":
    unittest.main()
