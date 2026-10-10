"""Split view for documents and spreadsheets: a chat's /work files shown beside the conversation.

GET /api/docview/{conv_id}/{path} answers with what the side panel (ui/src/components/DocPanel.tsx) draws:
- "doc"   Word files (.docx; .doc/.odt/.rtf through LibreOffice) as reflowing HTML built here from python-docx. Every
          piece of text is escaped and images are inlined as raster data URIs, so the markup carries nothing active.
- "sheet" Excel/Calc/CSV as a grid of display values per sheet, with formulas, bold cells and merges.
- "pdf"   PDFs as-is; slide decks converted to PDF (`?pdf=1` serves the converted file).
- "markdown" / "text" for plain text.

Big files are cut (rows, columns, characters) and say so with `truncated`. Results are cached by path, size and mtime.
"""
import asyncio
import base64
import csv
import html
import io
import re
from datetime import date, datetime, time as dtime
from pathlib import Path

from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse

import doc_tools
from tools import WORKSPACES

router = APIRouter(prefix="/api/docview")

DOC_EXT = {".docx", ".doc", ".odt", ".rtf"}
SHEET_EXT = {".xlsx", ".xlsm", ".xls", ".ods", ".csv", ".tsv"}
SLIDE_EXT = {".pptx", ".ppt", ".odp"}
TEXT_EXT = {".md", ".markdown", ".txt", ".log"}
MAX_BYTES = 40 * 1024 * 1024
MAX_ROWS, MAX_COLS = 2000, 80
MAX_TEXT = 400_000
MAX_IMAGE = 4 * 1024 * 1024
FORMULA_BYTES = 8 * 1024 * 1024    # formulas need a second workbook load; skipped above this
IMAGE_TYPES = {"image/png", "image/jpeg", "image/gif", "image/webp", "image/bmp"}

_cache: dict[str, dict] = {}


def _resolve(conv_id: str, path: str) -> Path:
    root = (WORKSPACES / conv_id).resolve()
    p = (root / path).resolve()
    if root.parent != WORKSPACES.resolve() or not str(p).startswith(str(root) + "/") or not p.is_file():
        raise HTTPException(404)
    return p


# ── documents ───────────────────────────────────────────────────────────────

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def _images(run_el, part) -> str:
    out = []
    for blip in run_el.iter(A_BLIP):
        rel = part.related_parts.get(blip.get(R_EMBED, ""))
        if rel is None or rel.content_type not in IMAGE_TYPES or len(rel.blob) > MAX_IMAGE:
            continue
        out.append(f'<img alt="" src="data:{rel.content_type};base64,{base64.b64encode(rel.blob).decode()}">')
    return "".join(out)


def _run_html(run_el, part) -> str:
    """One <w:r>: its text with bold/italic/underline/strike, tabs, breaks and pictures."""
    parts = []
    for ch in run_el:
        tag = ch.tag
        if tag in (W + "t", W + "delText"):
            parts.append(html.escape(ch.text or ""))
        elif tag == W + "tab":
            parts.append("\t")
        elif tag in (W + "br", W + "cr"):
            parts.append("<br>")
        elif tag == W + "drawing" or tag.endswith("}pict"):
            parts.append(_images(ch, part))
    text = "".join(parts)
    if not text:
        return ""
    rpr = run_el.find(W + "rPr")

    def on(name):
        el = rpr.find(W + name) if rpr is not None else None
        return el is not None and el.get(W + "val", "true") not in ("false", "0", "none")
    if on("b"):
        text = f"<strong>{text}</strong>"
    if on("i"):
        text = f"<em>{text}</em>"
    if on("u"):
        text = f"<u>{text}</u>"
    if on("strike"):
        text = f"<s>{text}</s>"
    va = rpr.find(W + "vertAlign") if rpr is not None else None
    if va is not None and va.get(W + "val") in ("superscript", "subscript"):
        t = "sup" if va.get(W + "val") == "superscript" else "sub"
        text = f"<{t}>{text}</{t}>"
    return text


def _inline_html(p_el, part) -> str:
    out = []
    for ch in p_el:
        if ch.tag == W + "r":
            out.append(_run_html(ch, part))
        elif ch.tag == W + "hyperlink":
            inner = "".join(_run_html(r, part) for r in ch.iter(W + "r"))
            rel = part.rels.get(ch.get(R_ID, ""))
            url = rel.target_ref if rel is not None and rel.is_external else ""
            if url.lower().startswith(("http://", "https://", "mailto:")):
                out.append(f'<a href="{html.escape(url)}" target="_blank" rel="noreferrer">{inner}</a>')
            else:
                out.append(inner)
        elif ch.tag in (W + "ins", W + "smartTag", W + "sdt", W + "fldSimple"):
            out.append("".join(_run_html(r, part) for r in ch.iter(W + "r")))
    return "".join(out)


def _para(p, part) -> tuple[str, str | None]:
    """(html, list kind or None) for one python-docx Paragraph."""
    style = (p.style.name if p.style is not None else "") or ""
    inner = _inline_html(p._p, part)
    ppr = p._p.pPr
    list_kind = None
    if (ppr is not None and ppr.numPr is not None) or style.startswith("List"):
        list_kind = "ol" if "Number" in style else "ul"
    align = ""
    if p.alignment is not None and int(p.alignment) in (1, 2, 3):
        align = f' style="text-align:{ {1: "center", 2: "right", 3: "justify"}[int(p.alignment)] }"'
    m = re.match(r"Heading (\d)", style)
    if style == "Title":
        return f"<h1{align}>{inner}</h1>", None
    if style == "Subtitle":
        return f'<p class="subtitle"{align}>{inner}</p>', None
    if m:
        n = min(int(m.group(1)) + 1, 6) if style != "Heading 1" else 2
        return f"<h{n}{align}>{inner}</h{n}>", None
    if list_kind:
        return f"<li>{inner}</li>", list_kind
    if style in ("Quote", "Intense Quote"):
        return f"<blockquote>{inner}</blockquote>", None
    return (f"<p{align}>{inner}</p>" if inner.strip() else "<p><br></p>"), None


def _table(t, part) -> str:
    from docx.table import _Cell
    rows = []
    for tr in t._tbl.tr_lst:
        cells = []
        for tc in tr.tc_lst:
            vm = tc.tcPr.vMerge if tc.tcPr is not None else None
            if vm is not None and vm.val != "restart":
                continue  # continuation of a vertically merged cell
            span = tc.grid_span
            body = _blocks(_Cell(tc, t), part)
            cells.append(f"<td{f' colspan={span}' if span > 1 else ''}>{body}</td>")
        rows.append(f"<tr>{''.join(cells)}</tr>")
    return f"<table>{''.join(rows)}</table>"


def _blocks(container, part) -> str:
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    out, open_list = [], None
    for item in container.iter_inner_content():
        if isinstance(item, Paragraph):
            h, kind = _para(item, part)
        elif isinstance(item, Table):
            h, kind = _table(item, part), None
        else:
            continue
        if kind != open_list:
            if open_list:
                out.append(f"</{open_list}>")
            if kind:
                out.append(f"<{kind}>")
            open_list = kind
        out.append(h)
    if open_list:
        out.append(f"</{open_list}>")
    return "".join(out)


def docx_html(p: Path) -> dict:
    import docx
    d = docx.Document(str(p))
    body = _blocks(d, d.part)
    truncated = len(body) > MAX_TEXT * 4
    if truncated:
        body = body[:MAX_TEXT * 4]
        body = body[:body.rfind("<")]
    return {"kind": "doc", "html": body, "truncated": truncated}


# ── spreadsheets ────────────────────────────────────────────────────────────

def _fmt(v, number_format: str = "General") -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == dtime(0) else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, (date, dtime)):
        return v.isoformat()
    if isinstance(v, (int, float)) and number_format != "General":
        fmt = number_format.split(";")[0]
        dec = fmt.split(".", 1)[1].count("0") if "." in fmt else 0
        if fmt.endswith("%"):
            return f"{v * 100:.{dec}f}%"
        if "0" in fmt:
            return f"{v:,.{dec}f}" if "," in fmt else f"{v:.{dec}f}"
    if isinstance(v, float):
        return f"{v:.10g}"
    return str(v)


def xlsx_grid(p: Path, values: Path | None = None) -> dict:
    """`values`: the same workbook recalculated by LibreOffice, for files whose formulas have no cached results
    (anything written by openpyxl rather than a spreadsheet app)."""
    import openpyxl
    wb = openpyxl.load_workbook(str(values or p), data_only=True)
    fwb = openpyxl.load_workbook(str(p), data_only=False) if p.stat().st_size <= FORMULA_BYTES else None
    sheets, uncached = [], False
    for ws in wb.worksheets:
        fws = fwb[ws.title] if fwb is not None else None
        nrows, ncols = ws.max_row or 0, ws.max_column or 0
        if nrows == 1 and ncols == 1 and ws.cell(1, 1).value is None:
            nrows = ncols = 0
        R, C = min(nrows, MAX_ROWS), min(ncols, MAX_COLS)
        if not R or not C:
            R = C = 0
        rows, formulas, bold = [], {}, []
        for r_i, row in enumerate(ws.iter_rows(min_row=1, max_row=R, max_col=C) if R else []):
            vals = []
            for c_i, cell in enumerate(row):
                vals.append(_fmt(cell.value, cell.number_format or "General"))
                if cell.font is not None and cell.font.b:
                    bold.append([r_i, c_i])
                if fws is not None:
                    f = fws.cell(r_i + 1, c_i + 1).value
                    if isinstance(f, str) and f.startswith("="):
                        formulas[f"{r_i},{c_i}"] = f
                        uncached = uncached or cell.value is None
            while len(vals) < C:
                vals.append("")
            rows.append(vals)
        merges = []
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= R and rng.min_col <= C:
                merges.append([rng.min_row - 1, rng.min_col - 1, min(rng.max_row, R) - 1, min(rng.max_col, C) - 1])
        widths = {}
        for key, dim in ws.column_dimensions.items():
            if dim.width and not dim.hidden:
                widths[key] = round(dim.width, 1)
        sheets.append({"name": ws.title, "rows": rows, "formulas": formulas, "bold": bold, "merges": merges,
                       "widths": widths, "nrows": nrows, "ncols": ncols, "truncated": nrows > R or ncols > C,
                       "hidden": ws.sheet_state != "visible"})
    return {"kind": "sheet", "sheets": sheets, "active": wb.index(wb.active) if wb.active in wb.worksheets else 0,
            "uncached": uncached and values is None}


def csv_grid(p: Path) -> dict:
    raw = p.read_bytes()[:MAX_BYTES]
    text = raw.decode("utf-8-sig", errors="replace")
    delim = "\t" if p.suffix.lower() == ".tsv" else ","
    if delim == ",":  # European CSVs use ";": go by the header line (csv.Sniffer gives up on ragged rows)
        head = text.split("\n", 1)[0]
        delim = max(",;\t|", key=head.count) if any(c in head for c in ",;\t|") else ","
    rows, nrows, ncols = [], 0, 0
    for row in csv.reader(io.StringIO(text), delimiter=delim):
        nrows += 1
        ncols = max(ncols, len(row))
        if nrows <= MAX_ROWS:
            rows.append(row[:MAX_COLS])
    C = min(ncols, MAX_COLS)
    rows = [r + [""] * (C - len(r)) for r in rows]
    return {"kind": "sheet", "active": 0, "sheets": [{
        "name": p.stem, "rows": rows, "formulas": {}, "bold": [], "merges": [], "widths": {},
        "nrows": nrows, "ncols": ncols, "truncated": nrows > MAX_ROWS or ncols > MAX_COLS, "hidden": False}]}


# ── endpoint ────────────────────────────────────────────────────────────────

async def build(p: Path, conv_id: str, path: str) -> dict:
    ext = p.suffix.lower()
    if p.stat().st_size > MAX_BYTES:
        raise HTTPException(413, "This file is too big for split view; download it instead.")
    try:
        if ext == ".docx":
            return await asyncio.to_thread(docx_html, p)
        if ext in DOC_EXT:
            return await asyncio.to_thread(docx_html, await doc_tools.soffice_convert(p, "docx"))
        if ext in (".xlsx", ".xlsm"):
            out = await asyncio.to_thread(xlsx_grid, p)
            if out.pop("uncached"):
                try:
                    out = await asyncio.to_thread(xlsx_grid, p, await doc_tools.soffice_convert(p, "xlsx"))
                    out.pop("uncached")
                except RuntimeError:
                    out["note"] = "Formulas show no results: this file was saved without them and LibreOffice isn't here to work them out."
            return out
        if ext in (".xls", ".ods"):
            out = await asyncio.to_thread(xlsx_grid, await doc_tools.soffice_convert(p, "xlsx"))
            out.pop("uncached")
            return out
        if ext in (".csv", ".tsv"):
            return await asyncio.to_thread(csv_grid, p)
        if ext == ".pdf":
            return {"kind": "pdf", "url": f"/api/workspace/{conv_id}/{path}"}
        if ext in SLIDE_EXT:
            await doc_tools.soffice_convert(p, "pdf")
            return {"kind": "pdf", "url": f"/api/docview/{conv_id}/{path}?pdf=1"}
        if ext in TEXT_EXT:
            text = p.read_text(errors="replace")
            return {"kind": "markdown" if ext in (".md", ".markdown") else "text",
                    "text": text[:MAX_TEXT], "truncated": len(text) > MAX_TEXT}
    except HTTPException:
        raise
    except RuntimeError as e:  # LibreOffice missing or failing
        raise HTTPException(422, str(e))
    except Exception as e:
        raise HTTPException(422, f"Couldn't read {p.name}: {type(e).__name__}: {e}"[:300])
    raise HTTPException(415, "Split view shows documents, spreadsheets, PDFs and text.")


@router.get("/{conv_id}/{path:path}")
async def view(conv_id: str, path: str, pdf: bool = False):
    p = _resolve(conv_id, path)
    if pdf:
        if p.suffix.lower() not in SLIDE_EXT:
            raise HTTPException(404)
        try:
            out = await doc_tools.soffice_convert(p, "pdf")
        except RuntimeError as e:
            raise HTTPException(422, str(e))
        return FileResponse(out, media_type="application/pdf", headers={"X-Content-Type-Options": "nosniff"})
    st = p.stat()
    key = f"{p}|{st.st_size}|{st.st_mtime_ns}"
    if key not in _cache:
        if len(_cache) > 32:
            _cache.pop(next(iter(_cache)))
        _cache[key] = await build(p, conv_id, path)
    return {**_cache[key], "name": p.name, "path": path, "size": st.st_size, "mtime": st.st_mtime}
